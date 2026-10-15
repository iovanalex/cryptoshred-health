"""Build the conference manuscript using the official IOS Press Word template."""
import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def no_numbering(paragraph):
    props = paragraph._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    value = OxmlElement("w:numId")
    value.set(qn("w:val"), "0")
    num.append(value)
    props.append(num)


def build(source, template, output):
    document = Document(template)
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                paragraph.clear()
    for style in document.styles:
        if style.type == 1:
            style.font.color.rgb = RGBColor(0, 0, 0)
            language = OxmlElement("w:lang")
            language.set(qn("w:val"), "en-GB")
            style._element.get_or_add_rPr().append(language)
    for name in ["Heading 1", "Heading 2", "HeadingUnn1"]:
        document.styles[name].paragraph_format.space_before = Pt(12)
        document.styles[name].paragraph_format.space_after = Pt(6)
        document.styles[name].paragraph_format.first_line_indent = Cm(0)
    document.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)

    blocks = source.read_text().strip().split("\n\n")
    abstract = False
    references = False
    first_body = True
    for index, block in enumerate(blocks):
        if block.startswith("# "):
            document.add_paragraph(block[2:], "Title")
        elif index == 1:
            document.add_paragraph(block, "Author")
        elif index == 2:
            document.add_paragraph(block, "Affiliation")
        elif index == 3:
            p = document.add_paragraph(block, "Affiliation")
            for r in p.runs:
                r.font.size = Pt(8)
        elif block == "## Abstract":
            abstract = True
        elif abstract:
            document.add_paragraph("Abstract. " + block, "Abstract")
            abstract = False
        elif block.startswith("Keywords:"):
            document.add_paragraph(block.replace("Keywords:", "Keywords.", 1), "Keywords")
        elif block.startswith("## "):
            references = block == "## References"
            document.add_paragraph(block[3:], "HeadingUnn1" if references else "Heading 1")
            first_body = True
        elif block.startswith("### "):
            document.add_paragraph(block[4:], "Heading 2")
            first_body = True
        elif block.startswith("Table "):
            p = document.add_paragraph(block, "CaptionLong")
            p.paragraph_format.keep_with_next = True
        elif block.startswith("|"):
            rows = [[c.strip() for c in row.strip().strip("|").split("|")]
                    for row in block.splitlines() if not re.fullmatch(r"[| :\-]+", row)]
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            widths = [0.7, 6.0, 5.7] if len(rows[0]) == 3 else [6.0, 1.6, 1.6, 1.6, 1.6]
            for col, width in zip(table.columns, widths):
                col.width = Cm(width)
            props = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                el = OxmlElement("w:" + edge)
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    el.set(qn("w:" + key), value)
                borders.append(el)
            props.append(borders)
            for row_index, values in enumerate(rows):
                row = table.add_row()
                trpr = row._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
                for column, (cell, value) in enumerate(zip(row.cells, values)):
                    cell.width = Cm(widths[column])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for side in ["top", "left", "bottom", "right"]:
                        el = OxmlElement("w:" + side)
                        el.set(qn("w:w"), "65")
                        el.set(qn("w:type"), "dxa")
                        margins.append(el)
                    tcpr.append(margins)
                    if row_index == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "EDEDED")
                        tcpr.append(shade)
                    p = cell.paragraphs[0]
                    p.style = "Table"
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (len(widths) == 3 and column > 0) or (len(widths) == 5 and column == 0) else WD_ALIGN_PARAGRAPH.CENTER
                    r = p.add_run(value)
                    r.font.size = Pt(8)
                    r.bold = row_index == 0
            document.add_paragraph().paragraph_format.space_after = Pt(3)
        elif references:
            p = document.add_paragraph(block, "References")
            no_numbering(p)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_after = Pt(2)
        else:
            document.add_paragraph(block, "NoindentNormal" if first_body else "Normal")
            first_body = False
    document.core_properties.title = blocks[0][2:]
    document.core_properties.author = "Robert Hevesi; Alexandru Iovanovici"
    document.core_properties.subject = "Scenario-based EHR erasure evaluation for MIE 2027"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--source", type=Path, default=Path(__file__).with_name("manuscript.md"))
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.template, args.output)
